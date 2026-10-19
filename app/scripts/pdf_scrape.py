import io
import logging
import os
import uuid
from urllib.parse import urlparse
import cv2
import pymupdf as fitz
import numpy as np
import pytesseract
from PIL import Image
import chromadb
import docx
import openpyxl
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
# from sentence_transformers import SentenceTransformer

from app import data_access as S3_DATA
from app import constant as CONSTANT

logger = logging.getLogger(__name__)

admin_client = chromadb.AdminClient()

try:
    admin_client.get_tenant(name=CONSTANT.CHROMA_TENANT)
except Exception:
    admin_client.create_tenant(name=CONSTANT.CHROMA_TENANT)

try:
    admin_client.get_database(
        name=CONSTANT.CHROMA_DB,
        tenant=CONSTANT.CHROMA_TENANT
    )
except Exception:
    admin_client.create_database(
        name=CONSTANT.CHROMA_DB,
        tenant=CONSTANT.CHROMA_TENANT
    )

class PDFExtractor:
    def __init__(self, directory_path):
        """Initializes with the folder path containing the files."""
        self.directory_path = directory_path

    def extract_all(self):
        """Step 1: Get the extracted text from all files in the path."""
        all_extracted_data = {}

        if self.directory_path.startswith("s3://"):
            s3_client = S3_DATA.s3_client
            parsed_url = urlparse(self.directory_path)
            bucket_name = parsed_url.netloc
            prefix = parsed_url.path.lstrip("/")

            try:
                response = s3_client.list_objects_v2(
                    Bucket=bucket_name, Prefix=prefix
                )
                if "Contents" not in response:
                    logger.warning(f"No files found in S3 path: {self.directory_path}")
                    return all_extracted_data

                for obj in response["Contents"]:
                    file_key = obj["Key"]
                    if file_key == prefix or file_key.endswith("/"):
                        continue

                    file_name = os.path.basename(file_key)
                    logger.info(f"Processing file from S3: {file_name}")

                    file_obj = s3_client.get_object(Bucket=bucket_name, Key=file_key)
                    file_bytes = io.BytesIO(file_obj["Body"].read())

                    if file_name.lower().endswith(".pdf"):
                        all_extracted_data[file_name] = self._extract_pdf(file_bytes)
                    elif file_name.lower().endswith(".docx"):
                        all_extracted_data[file_name] = self._extract_docx(file_bytes)
                    elif file_name.lower().endswith(".xlsx"):
                        all_extracted_data[file_name] = self._extract_xlsx(file_bytes)
                    else:
                        logger.warning(f"Unsupported file format: {file_name}")

                return all_extracted_data

            except Exception as e:
                logger.error(f"Failed to process S3 directory: {e}")
                return all_extracted_data

        return all_extracted_data

    def _extract_pdf(self, file_path):
        page_text = {}
        image_text = {}
        stream = file_path.read() if isinstance(file_path, io.BytesIO) else file_path
        with fitz.open(stream=stream, filetype="pdf") as pdf:
            for page_num, page in enumerate(pdf, start=1):
                page_text[page_num] = page.get_text()
                image_text[page_num] = []
                image_list = page.get_images(full=True)
                for img_info in image_list:
                    try:
                        xref = img_info[0]
                        base_image = pdf.extract_image(xref)
                        image_bytes = base_image["image"]
                        image = Image.open(io.BytesIO(image_bytes)).convert("RGB")
                        img = np.array(image)
                        gray = cv2.cvtColor(img, cv2.COLOR_RGB2GRAY)
                        _, thresh = cv2.threshold(gray, 150, 255, cv2.THRESH_BINARY)
                        text = pytesseract.image_to_string(thresh)
                        image_text[page_num].append(text)
                    except Exception as img_err:
                        logger.warning(f"Failed to process an image on page {page_num}: {img_err}")

        extracted_text = {}
        for page_num in page_text:
            img_text = "\n".join(image_text[page_num])
            extracted_text[f"page_{page_num}"] = (page_text[page_num] + "\n" + img_text).strip()
        return extracted_text

    def _extract_docx(self, file_path):
        doc = docx.Document(file_path)
        content = []
        for para in doc.paragraphs:
            if para.text.strip():
                content.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    content.append(row_text)
        return {"content": "\n".join(content)}

    def _extract_xlsx(self, file_path):
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet_data = {}
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows = []
            for row in sheet.iter_rows(values_only=True):
                row_values = [str(cell) if cell is not None else "" for cell in row]
                if any(row_values):
                    rows.append("\t".join(row_values))
            sheet_data[sheet_name] = "\n".join(rows)
        return sheet_data

    def store_to_chroma(self, persist_directory=CONSTANT.CHROMA_DIRS, tenant=CONSTANT.CHROMA_TENANT, database=CONSTANT.CHROMA_DB):
        """Function to process and store text into ChromaDB."""
        extracted_data = self.extract_all()
        if not extracted_data:
            logger.warning("No data extracted from files.")
            return None

        documents = []
        for file_name, file_content in extracted_data.items():
            if isinstance(file_content, dict):
                for section_key, text_content in file_content.items():
                    if text_content.strip():
                        doc = Document(
                            page_content=text_content,
                            metadata={"source": file_name, "section": section_key, "extracted_data": text_content ,"s3_key": file_name,"document_name": os.path.basename(file_name)})
                        documents.append(doc)
                        
        text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=4000, 
            chunk_overlap=400
        )
        chunks = text_splitter.split_documents(documents)
        logger.info(f"Created {len(chunks)} text chunks.")

        client = chromadb.PersistentClient(
            path=persist_directory,
            settings=chromadb.Settings(tenant=tenant, database=database)
        )
        collection_name = CONSTANT.CHROMA_COLLECTIONS
 
        # try:
        #     client.delete_collection(name=collection_name)
        # except Exception:
        #     pass

        collection = client.get_or_create_collection(name=collection_name)
        logger.info(f"Using Tenant: '{tenant}' | Database: '{database}' | Collection: '{collection_name}' (Count: {collection.count()})")

        embeddings = HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2")

        ids = [str(uuid.uuid4()) for _ in chunks]
        texts = [chunk.page_content for chunk in chunks]
        metadatas = [chunk.doc.metadata for chunk in chunks]

        logger.info("Generating embeddings for chunks...")
        embedding_vectors = embeddings.embed_documents(texts)
        collection.add(
            ids=ids,
            documents=texts,
            metadatas=metadatas,
            embeddings=embedding_vectors
        )
        logger.info(f"ids: {len(ids)} | documents: {len(documents)} | metadatas: {len(metadatas)}")
        
        logger.info(f"Successfully embedded and stored {len(chunks)} chunks into ChromaDB collection at '{persist_directory}'!")
        return collection

if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    bucket_name = S3_DATA.BUCKET_NAME 
    s3_directory = S3_DATA.S3_DIRECTORY.strip("/")
    s3_folder_path = f"s3://{bucket_name}/{s3_directory}/"
    extractor = PDFExtractor(s3_folder_path)
    extractor.store_to_chroma()